import pptx

prs = pptx.Presentation('SMART INDIA HACKATHON 2025 Sikkim (2).pptx')
for idx in range(len(prs.slides)):
    slide = prs.slides[idx]
    print(f"=== SLIDE {idx+1} ===")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip().replace('\n', ' // ')
            if txt:
                print(f"  Shape {s_idx} ({shape.name}): {txt[:80]}")
