import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) of a table cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    Set cell borders.
    kwargs can be top, bottom, left, right.
    val: 'single', 'double', 'dashed', etc.
    color: '1A365D', 'CCCCCC', etc.
    sz: '4', '8', '12', etc. (1/8 pt)
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    tcBorders = tc_pr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tc_pr.append(tcBorders)
    
    for edge, border_args in kwargs.items():
        edge_data = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{border_args.get("val", "single")}" '
            f'w:sz="{border_args.get("sz", "4")}" w:space="0" '
            f'w:color="{border_args.get("color", "CCCCCC")}"/>'
        )
        tcBorders.append(edge_data)

def add_callout(doc, title, text, box_type="NOTE", color_hex="1A365D", bg_hex="F0F4F8"):
    """Adds a stylish callout box with a colored left border and shaded background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    # Left thick border, clear others
    set_cell_border(cell, 
                    left={"val": "single", "sz": "24", "color": color_hex},
                    top={"val": "none", "sz": "0", "color": "auto"},
                    bottom={"val": "none", "sz": "0", "color": "auto"},
                    right={"val": "none", "sz": "0", "color": "auto"})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"[{box_type.upper()}] {title}\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(int(color_hex[:2], 16), int(color_hex[2:4], 16), int(color_hex[4:], 16))
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.italic = (box_type.upper() == "TALKING POINT")
    run_text.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    # Add small spacing after callout
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_diagram_box(doc, title, diagram_text):
    """Adds a cleanly styled diagram box with monospaced code-like appearance."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=160, bottom=160, left=180, right=180)
    set_cell_border(cell, 
                    left={"val": "single", "sz": "8", "color": "CBD5E1"},
                    top={"val": "single", "sz": "8", "color": "CBD5E1"},
                    bottom={"val": "single", "sz": "8", "color": "CBD5E1"},
                    right={"val": "single", "sz": "8", "color": "CBD5E1"})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    
    run_title = p.add_run(f"ARCHITECTURE / WORKFLOW DIAGRAM: {title}\n\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(9.5)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    run_diag = p.add_run(diagram_text)
    run_diag.font.name = "Consolas"
    run_diag.font.size = Pt(8.5)
    run_diag.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def format_table(tbl, col_widths, headers, data, header_bg="1A365D", zebra=True):
    """Formats a table with a stylish dark header, zebra striping, and cell borders."""
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Headers
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        set_cell_border(hdr_cells[i], 
                        top={"val": "single", "sz": "6", "color": "0F172A"},
                        bottom={"val": "single", "sz": "12", "color": "0F172A"},
                        left={"val": "single", "sz": "4", "color": "334155"},
                        right={"val": "single", "sz": "4", "color": "334155"})
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Rows
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_color = "F8FAFC" if (zebra and row_idx % 2 == 1) else "FFFFFF"
        for col_idx, text_val in enumerate(row_data):
            row_cells[col_idx].text = str(text_val)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=90, bottom=90, left=110, right=110)
            set_cell_border(row_cells[col_idx], 
                            top={"val": "single", "sz": "4", "color": "E2E8F0"},
                            bottom={"val": "single", "sz": "4", "color": "E2E8F0"},
                            left={"val": "single", "sz": "4", "color": "E2E8F0"},
                            right={"val": "single", "sz": "4", "color": "E2E8F0"})
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                
    # Column widths
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

print("Helper functions defined successfully.")
